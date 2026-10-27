"""The rooming list as files for the hotel: Excel (what the producer sends today, and
the hotel can edit) and a PDF for printing. Both read one sheet from
services/rooming.build_sheets and nothing else, so neither can print a field
the sheet doesn't carry — no phones, emails, rates or per diem.

Layout follows the producer's own sheet (docs/producer-documents-map.md §1.5): a title,
Group and Contact lines, a date row, a weekday row and a day-type row over the
night columns, then Name · Notes · Reservation # · one column per date ·
TOTAL, and a TOTALS row (rooms a night). With extra rooms on the list (a merch
room) a Hotel Breakfast Count row follows — people only, as on her sheet.
"""

import io
from datetime import date
from pathlib import Path

from jinja2 import Environment, FileSystemLoader, select_autoescape

from services.job_dates import LABELS

_WEEKDAYS = ("Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun")


def _date_head(iso: str) -> tuple[str, str]:
    d = date.fromisoformat(iso)
    return f"{d.month}/{d.day}", _WEEKDAYS[d.weekday()]


def _type_head(col: dict) -> str:
    return " / ".join(LABELS[t] for t in col["types"])


def title(job: dict, sheet: dict) -> str:
    hotel = sheet.get("hotel") or {}
    return f"Hotel Rooming List - {hotel.get('name') or job.get('name') or 'Job'}"


def _header_lines(job: dict, sheet: dict) -> list[tuple[str, str]]:
    hotel = sheet.get("hotel") or {}
    lines = [("Group", job.get("client_name") or job.get("name") or "")]
    if hotel.get("contact"):
        lines.append(("Contact", hotel["contact"]))
    if hotel.get("check_in"):
        a, b = _date_head(hotel["check_in"]), _date_head(hotel["check_out"])
        lines.append(("Stay", f"{a[1]} {a[0]} – {b[1]} {b[0]}"))
    return lines


# ── Excel ────────────────────────────────────────────────────────────────────

def to_xlsx(job: dict, sheet: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Rooming list"
    bold = Font(bold=True)
    ws.append([title(job, sheet)])
    ws["A1"].font = Font(bold=True, size=14)
    for label, value in _header_lines(job, sheet):
        ws.append([f"{label}:", value])
        ws.cell(row=ws.max_row, column=1).font = bold
    ws.append([])

    cols = sheet["dates"]
    first = 4                                   # Name, Notes, Reservation #, then dates
    ws.append(["", "", "Dates"] + [_date_head(c["date"])[0] for c in cols] + [""])
    ws.append(["", "", "Days"] + [_date_head(c["date"])[1] for c in cols] + [""])
    ws.append(["", "", ""] + [_type_head(c) for c in cols] + [""])
    ws.append(["Crew Names", "Notes", "Reservation Number"] + ["" for _ in cols] + ["TOTAL"])
    for r in range(ws.max_row - 3, ws.max_row + 1):
        for c in range(1, first + len(cols) + 1):
            ws.cell(row=r, column=c).font = bold
            ws.cell(row=r, column=c).alignment = Alignment(horizontal="center" if c >= first else "left")

    top = ws.max_row + 1
    for row in sheet["rows"]:
        cells = []
        for c in cols:
            v = row["cells"].get(c["date"])
            cells.append(1 if v == "1" else (v or ""))
        ws.append([row["name"], row["notes"], row["reservation"]] + cells + [None])
        n = ws.max_row
        a = ws.cell(row=n, column=first).coordinate
        b = ws.cell(row=n, column=first + len(cols) - 1).coordinate
        ws.cell(row=n, column=first + len(cols), value=f"=SUM({a}:{b})")   # live, as in her sheet
    bottom = ws.max_row

    ws.append(["TOTALS", "", ""] + [None for _ in cols] + [None])
    n = ws.max_row
    for i in range(len(cols) + 1):
        col = ws.cell(row=n, column=first + i).column_letter
        ws.cell(row=n, column=first + i, value=f"=SUM({col}{top}:{col}{bottom})" if bottom >= top else 0)
        ws.cell(row=n, column=first + i).font = bold
    ws.cell(row=n, column=1).font = bold

    if any(r.get("kind") == "room" for r in sheet["rows"]):
        ws.append(["Hotel Breakfast Count", "", ""] + [sheet["breakfasts"].get(c["date"], 0) for c in cols] + [""])
        ws.cell(row=ws.max_row, column=1).font = bold

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 20
    for i in range(len(cols) + 1):
        ws.column_dimensions[ws.cell(row=1, column=first + i).column_letter].width = 11
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── PDF (HTML for WeasyPrint) ────────────────────────────────────────────────

_env: Environment | None = None


def _environment() -> Environment:
    global _env
    if _env is None:
        _env = Environment(loader=FileSystemLoader(Path(__file__).parent / "templates"),
                           autoescape=select_autoescape(["html", "j2", "html.j2"], default=True),
                           trim_blocks=True, lstrip_blocks=True)
    return _env


def to_html(job: dict, sheet: dict) -> str:
    cols = [{"date": _date_head(c["date"])[0], "weekday": _date_head(c["date"])[1],
             "type": _type_head(c), "key": c["date"]} for c in sheet["dates"]]
    return _environment().get_template("rooming_list.html.j2").render(
        title=title(job, sheet), header=_header_lines(job, sheet), cols=cols,
        rows=sheet["rows"], totals=sheet["totals"], total=sheet["total"],
        breakfasts=sheet["breakfasts"] if any(r.get("kind") == "room" for r in sheet["rows"]) else None)
