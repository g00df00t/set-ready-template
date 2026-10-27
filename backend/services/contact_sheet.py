"""The job's contact sheet — the top half of the producer's Job Info Sheet
(docs/producer-documents-map.md §1.9), which she "constantly" fills by copying
and pasting crew and talent details in and out. Here it's a VIEW: nothing is
typed twice, every row reads where the person already lives. Pure — the
router fetches, this lays it out, and the Excel copy reads only this.

Sections, in her order: crew (with their "travel to travel" dates), talent
(with agency and agent), client & agency contacts from the job's days (one row
per person, however many days they're on), and vendor contacts from the
Vendors tab. Released hires are left off; holds show as holds.

⚠️ Producer-only. It has everyone's phone and email, so it's never shared,
emailed or printed on a crew-facing document.
"""

import io

from services.holds import status_label, status_of
from services.job_dates import hire_base
from services.rooming import their_dates


def _span(dates: list[str]) -> tuple[str | None, str | None]:
    return (dates[0], dates[-1]) if dates else (None, None)


def _clean(v) -> str:
    return (v or "").strip()


def build(job: dict, hires: list[dict], roster_by_id: dict, day_contacts: list[dict],
          vendors: list[dict]) -> dict:
    crew, talent = [], []
    for h in sorted(hires, key=lambda h: h.get("sort_order") or 0):
        if status_of(h) == "released":
            continue
        first, last = _span(their_dates(hire_base(job, h), h.get("dates_override")))
        row = {"hire_id": h["id"], "name": h.get("name") or "", "role": _clean(h.get("role")),
               "phone": _clean(h.get("phone")), "email": _clean(h.get("email")),
               "first": first, "last": last,
               "status": status_label(h) if status_of(h) == "hold" else ""}
        if h.get("kind") == "talent":
            r = roster_by_id.get(str(h.get("roster_id"))) or {}
            talent.append({**row, "segment": _clean(h.get("segment")), "agency": _clean(r.get("agency")),
                           "agent": _clean(r.get("agent_name")), "agent_phone": _clean(r.get("agent_phone")),
                           "agent_email": _clean(r.get("agent_email"))})
        else:
            crew.append(row)

    # One row per person across the job's days: matched by email, else by
    # name and organisation. The first non-empty value of each field wins.
    clients: dict[tuple, dict] = {}
    for c in day_contacts:
        email = _clean(c.get("email")).lower()
        key = ("e", email) if email else ("n", _clean(c.get("name")).lower(), _clean(c.get("organization")).lower())
        row = clients.setdefault(key, {"name": "", "organization": "", "role": "", "phone": "", "email": ""})
        for f in ("name", "organization", "role", "phone", "email"):
            if not row[f]:
                row[f] = _clean(c.get(f))
    client_rows = sorted(clients.values(), key=lambda r: (r["organization"].lower(), r["name"].lower()))

    vendor_rows = [{"kind": v["kind_label"], "title": v["title"], "name": _clean(v.get("contact_name")),
                    "phone": _clean(v.get("contact_phone")), "email": _clean(v.get("contact_email"))}
                   for v in vendors
                   if not v.get("broken") and (v.get("contact_name") or v.get("contact_phone") or v.get("contact_email"))]
    return {"crew": crew, "talent": talent, "clients": client_rows, "vendors": vendor_rows}


def title(job: dict) -> str:
    return f"Job Info - {job.get('name') or 'Job'}"


def to_xlsx(job: dict, sheet: dict, usage: list[tuple[str, str]]) -> bytes:
    """Her sheet's top half as a spreadsheet she can paste from."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Contacts"
    bold = Font(bold=True)
    ws.append([title(job)])
    ws["A1"].font = Font(bold=True, size=14)
    for label, value in usage:
        ws.append([label, value])
        ws.cell(row=ws.max_row, column=1).font = bold

    def block(heading: str, columns: list[str], rows: list[list]):
        ws.append([])
        ws.append([heading])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
        ws.append(columns)
        for c in range(1, len(columns) + 1):
            ws.cell(row=ws.max_row, column=c).font = bold
        for r in rows:
            ws.append(r)
        if not rows:
            ws.append(["—"])

    def dates(r):
        if not r["first"]:
            return ""
        return r["first"] if r["first"] == r["last"] else f"{r['first']} to {r['last']}"

    block("CREW", ["Role", "Name", "Cell", "Email", "Travel to Travel Dates", "Status"],
          [[r["role"], r["name"], r["phone"], r["email"], dates(r), r["status"]] for r in sheet["crew"]])
    block("TALENT", ["Segment", "Name", "Cell", "Email", "Travel to Travel Dates", "Agency", "Agent", "Agent Cell",
                     "Agent Email", "Status"],
          [[r["segment"], r["name"], r["phone"], r["email"], dates(r), r["agency"], r["agent"], r["agent_phone"],
            r["agent_email"], r["status"]] for r in sheet["talent"]])
    block("CLIENT & AGENCY", ["Organization", "Name", "Role", "Cell", "Email"],
          [[r["organization"], r["name"], r["role"], r["phone"], r["email"]] for r in sheet["clients"]])
    block("VENDORS", ["Kind", "Vendor", "Contact", "Cell", "Email"],
          [[r["kind"], r["title"], r["name"], r["phone"], r["email"]] for r in sheet["vendors"]])
    for col, w in zip("ABCDEFGHIJ", (18, 24, 16, 28, 24, 20, 20, 16, 28, 12)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
