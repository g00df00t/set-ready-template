"""The travel doc and the car service's copy, as Excel and as a landscape PDF.
Both read services/travel.build's rows and nothing else.

Layout follows the producer's Travel Doc (docs/producer-documents-map.md §1.7): an
ARRIVALS block, then DEPARTURES, columns Name · Date · Airline · Flight # ·
Depart · Dept Time · Arrival Date · Arrival · Arrival Time · Transportation
Plan.

    travel doc    every booked person's legs. NO phone column — it's hers
    car service   only legs planned as car service, with a Phone column: the
                  person's own, or the producer's for anyone switched to it
                  (marked "(producer)"). Nothing else personal: no email, no
                  role detail, no rate
"""

import io
from datetime import date
from pathlib import Path

from jinja2 import Environment, FileSystemLoader, select_autoescape

from services.travel import display_time

_WEEKDAYS = ("Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun")

BASE_COLUMNS = ["Name", "Date", "Airline", "Flight #", "Depart", "Dept Time",
                "Arrival Date", "Arrival", "Arrival Time", "Transportation Plan"]
# The PDF's column widths (%), so ARRIVALS and DEPARTURES line up.
_WIDTHS = [14, 8, 8, 8, 7, 8, 9, 7, 8, 23]
_PHONE_WIDTH = 12


def _day(iso: str | None) -> str:
    if not iso:
        return ""
    d = date.fromisoformat(iso)
    return f"{_WEEKDAYS[d.weekday()]} {d.month}/{d.day}"


def title(job: dict, car_service: bool) -> str:
    what = "Car Service" if car_service else "Travel"
    return f"{what} - {job.get('name') or 'Job'}"


def _header(job: dict, view: dict, car_service: bool) -> list[tuple[str, str]]:
    lines = []
    if job.get("client_name"):
        lines.append(("Group", job["client_name"]))
    if car_service and view.get("producer_phone"):
        lines.append(("Production contact", view["producer_phone"]))
    return lines


def _cells(r: dict, car_service: bool) -> list[str]:
    cells = [r["name"], _day(r["date"]), r["carrier"], r["flight"], r["from_airport"],
             display_time(r["depart_time"]), _day(r["arrive_date"]), r["to_airport"],
             display_time(r["arrive_time"]), r["plan_text"]]
    if car_service:
        phone = r.get("phone") or ""
        cells.insert(1, f"{phone} (producer)" if phone and r.get("phone_is_producer") else phone)
    return cells


def blocks(view: dict, car_service: bool) -> tuple[list[str], list[tuple[str, list[list[str]]]]]:
    """(the column heads, [(ARRIVALS, rows), (DEPARTURES, rows)])."""
    columns = list(BASE_COLUMNS)
    if car_service:
        columns.insert(1, "Phone")
    source = view["manifest"] if car_service else view["arrivals"] + view["departures"]
    out = []
    for label, direction in (("ARRIVALS", "in"), ("DEPARTURES", "out")):
        out.append((label, [_cells(r, car_service) for r in source if r["direction"] == direction]))
    return columns, out


def to_xlsx(job: dict, view: dict, car_service: bool) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Car service" if car_service else "Travel"
    bold = Font(bold=True)
    ws.append([title(job, car_service)])
    ws["A1"].font = Font(bold=True, size=14)
    for label, value in _header(job, view, car_service):
        ws.append([f"{label}:", value])
        ws.cell(row=ws.max_row, column=1).font = bold

    columns, sections = blocks(view, car_service)
    for label, rows in sections:
        ws.append([])
        ws.append([label])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
        ws.append(columns)
        for c in range(1, len(columns) + 1):
            ws.cell(row=ws.max_row, column=c).font = bold
        for row in rows:
            ws.append(row)
            for c in range(1, len(columns) + 1):
                ws.cell(row=ws.max_row, column=c).alignment = Alignment(vertical="top", wrap_text=True)
        if not rows:
            ws.append(["—"])

    widths = [24] + ([22] if car_service else []) + [12, 14, 11, 9, 11, 12, 9, 11, 34]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_env: Environment | None = None


def _environment() -> Environment:
    global _env
    if _env is None:
        _env = Environment(loader=FileSystemLoader(Path(__file__).parent / "templates"),
                           autoescape=select_autoescape(["html", "j2", "html.j2"], default=True),
                           trim_blocks=True, lstrip_blocks=True)
    return _env


def to_html(job: dict, view: dict, car_service: bool) -> str:
    columns, sections = blocks(view, car_service)
    widths = list(_WIDTHS)
    if car_service:
        widths.insert(1, _PHONE_WIDTH)
        widths[-1] -= _PHONE_WIDTH
    return _environment().get_template("travel_doc.html.j2").render(
        title=title(job, car_service), header=_header(job, view, car_service),
        columns=columns, widths=widths, sections=sections)
