"""The talent hold sheet (docs/producer-documents-map.md §1.2): every talent on
the job, grouped into her segments ("Brand Women", "Brand Men"), with their
agency, agent, rate plus the agency's fee, home city, notes, portfolio link —
and 1 / 2 / ✓ across the job's dates. Pure: the router fetches the hires and
their Roster rows, this lays them out; rooming_documents-style files follow.

Booked is "X", not ✓: the server's PDF fonts printed ✓ blank on staging
(2026-10-02), and an X reads the same in every font, Excel included. The
Holds page in the browser keeps its ✓. LEGEND says which is which.

Rates: the hire's agreed shoot-day rate (services/job_rates.py), printed with
the agency fee on top the way she writes it: "$1,500 + 20%". Never the fee
worked out as a total — the sheet is for holding, not invoicing.
"""

import io
from datetime import date
from pathlib import Path

from jinja2 import Environment, FileSystemLoader, select_autoescape

from services.holds import _job_dates, status_of
from services.job_dates import DAY_TYPES, LABELS, effective_dates, hire_base, works_on
from services.job_rates import _num, money

_WEEKDAYS = ("Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun")
NO_SEGMENT = "Talent"
BOOKED = "X"
LEGEND = "1 = first hold · 2 = second hold · X = booked"


def rate_text(hire: dict) -> str:
    """"$1,500 + 20%"; just the rate with no fee; "" with no rate."""
    rate = _num(hire.get("rate_shoot"))
    if rate is None:
        return ""
    fee = _num(hire.get("agency_fee_pct"))
    return f"{money(rate)} + {fee:g}%" if fee else money(rate)


def safe_link(url) -> str:
    """Only a web address becomes a link — checked again here, at the point it
    renders, whatever path put it on the Roster."""
    url = (url or "").strip()
    return url if url.lower().startswith(("http://", "https://")) else ""


def agent_text(roster: dict) -> str:
    """Her Agent column holds a name, and sometimes the email."""
    return roster.get("agent_name") or roster.get("agent_email") or ""


def build_talent_sheet(job: dict, hires: list[dict], roster_by_id: dict) -> dict:
    dates = _job_dates(job)
    crew = job.get("crew_dates") or {}
    segments: dict[str, list] = {}
    for h in sorted(hires, key=lambda h: h.get("sort_order") or 0):
        if h.get("kind") != "talent" or status_of(h) == "released":
            continue
        r = roster_by_id.get(str(h.get("roster_id"))) or {}
        eff = effective_dates(hire_base(job, h), h.get("dates_override"))
        mark = BOOKED if status_of(h) == "confirmed" else str(h.get("hold_rank") or 1)
        segments.setdefault(h.get("segment") or NO_SEGMENT, []).append({
            "hire_id": h["id"], "name": h.get("name"),
            "agency": r.get("agency") or "", "agent": agent_text(r),
            "rate": rate_text(h), "based": r.get("based") or "",
            "notes": h.get("notes") or "", "link": safe_link(r.get("portfolio_url")),
            "cells": {d: mark for d in dates if works_on(eff, d)},
        })
    heads = []
    for d in dates:
        types = [t for t in DAY_TYPES if any(a <= d <= b for a, b in crew.get(t, []))]
        heads.append({"date": d, "what": " / ".join(LABELS[t] for t in types).upper()})
    return {"title": " – ".join(p for p in (job.get("client_name"), job.get("name")) if p),
            "dates": heads, "segments": [{"name": k, "rows": v} for k, v in segments.items()]}


def _date_head(iso: str) -> tuple[str, str]:
    d = date.fromisoformat(iso)
    return f"{d.month}/{d.day}", _WEEKDAYS[d.weekday()]


COLUMNS = ("AGENCY", "AGENT", "TALENT", "RATE + FEE", "BASED", "NOTES", "Link")


def to_xlsx(sheet: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Talent holds"
    bold = Font(bold=True)
    ws.append([sheet["title"] or "Talent holds"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([LEGEND])
    ws.append([])
    first = len(COLUMNS) + 1
    for seg in sheet["segments"]:
        # Each segment repeats the header block, as hers does.
        ws.append([seg["name"]])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
        ws.append([""] * len(COLUMNS) + [_date_head(c["date"])[0] for c in sheet["dates"]])
        ws.append([""] * len(COLUMNS) + [_date_head(c["date"])[1] for c in sheet["dates"]])
        ws.append(list(COLUMNS) + [c["what"] for c in sheet["dates"]])
        for r in range(ws.max_row - 2, ws.max_row + 1):
            for c in range(1, first + len(sheet["dates"])):
                ws.cell(row=r, column=c).font = bold
                ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor="F2F2F2")
        for row in seg["rows"]:
            ws.append([row["agency"], row["agent"], row["name"], row["rate"], row["based"], row["notes"], row["link"]]
                      + [row["cells"].get(c["date"], "") for c in sheet["dates"]])
            if row["link"]:
                cell = ws.cell(row=ws.max_row, column=7)
                cell.hyperlink = row["link"]
                cell.font = Font(color="1F5FB0", underline="single")
        ws.append([])
    for col, width in zip("ABCDEFG", (18, 20, 22, 16, 16, 28, 26)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_env: Environment | None = None


def to_html(sheet: dict) -> str:
    global _env
    if _env is None:
        _env = Environment(loader=FileSystemLoader(Path(__file__).parent / "templates"),
                           autoescape=select_autoescape(["html", "j2", "html.j2"], default=True),
                           trim_blocks=True, lstrip_blocks=True)
    cols = [{"key": c["date"], "date": _date_head(c["date"])[0], "weekday": _date_head(c["date"])[1], "what": c["what"]}
            for c in sheet["dates"]]
    return _env.get_template("talent_hold_sheet.html.j2").render(title=sheet["title"] or "Talent holds",
                                                                  segments=sheet["segments"], cols=cols, legend=LEGEND)
