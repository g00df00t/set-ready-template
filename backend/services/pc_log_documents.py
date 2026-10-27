"""The PC log as files: Excel in the producer's own layout (PC log template)
and a landscape PDF. Both read one log — the live one, or the copy frozen when
it was marked sent — and nothing else.

Her sheet: JOB NAME / PRODUCER / PROJECT DATES at the top left with the
category legend beside them; then # · DATE · TO WHOM PAID · FOR WHAT · TOTAL ·
the nine columns; then TOTAL FEES AND EXPENSES, PRODUCTION FEE %, GRAND TOTAL,
LESS ADVANCE RECEIVED, DIFFERENCE. In Excel every total is a live formula, so
she can still change a cell and have the footer follow.

⚠️ The PDF is ASCII marks only: no dashes, dots or ticks outside plain ASCII
in anything this module writes (what the producer typed prints as typed).
"""

import io
from datetime import date
from pathlib import Path

from jinja2 import Environment, FileSystemLoader, select_autoescape

from services.pc_log import CATEGORIES, LEGEND, dollars

MONEY_FORMAT = '"$"#,##0.00'


def title(log: dict) -> str:
    return f"PC Log - {log['header'].get('job_name') or 'Job'}"


def _date(iso: str | None) -> str:
    if not iso:
        return ""
    d = date.fromisoformat(str(iso))
    return f"{d.month}/{d.day}/{d.year % 100:02d}"


def shown_rows(log: dict) -> list[dict]:
    return [r for r in log["rows"] if not r.get("hidden")]


def to_xlsx(log: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "PC Log"
    bold = Font(bold=True)
    h = log["header"]
    cats = list(CATEGORIES)

    for row, (label, value) in enumerate((("JOB NAME:", h.get("job_name")), ("PRODUCER:", h.get("producer")),
                                          ("PROJECT DATES:", h.get("project_dates"))), start=1):
        ws.cell(row=row, column=1, value=label).font = bold
        ws.cell(row=row, column=2, value=value or "")
    # The legend, in two columns beside the header, as on her sheet.
    for i, k in enumerate(cats):
        col, row = (5, i + 1) if i < 5 else (10, i - 4)
        ws.cell(row=row, column=col, value=f"{CATEGORIES[k]} - {LEGEND[k]}")
    fee_pct = log["totals"]["fee_bp"] / 100
    ws.cell(row=5, column=10, value=f"Production Fee - {fee_pct:g}% of fees and expenses")
    ws.cell(row=6, column=10, value=f"{CATEGORIES['payroll_fee']} - {LEGEND['payroll_fee']}")

    heads = ["#", "DATE", "TO WHOM PAID", "FOR WHAT", "TOTAL"] + [CATEGORIES[k] for k in cats]
    for c, text in enumerate(heads, start=1):
        cell = ws.cell(row=8, column=c, value=text)
        cell.font = bold
        cell.alignment = Alignment(horizontal="center" if c >= 5 else "left", wrap_text=True)

    first = r = 9
    for line in shown_rows(log):
        ws.cell(row=r, column=1, value=line.get("number"))
        ws.cell(row=r, column=2, value=_date(line.get("spent_on")))
        ws.cell(row=r, column=3, value=line.get("payee") or "")
        ws.cell(row=r, column=4, value=line.get("description") or "")
        col = 6 + cats.index(line["category"])
        ws.cell(row=r, column=col, value=line["amount_cents"] / 100).number_format = MONEY_FORMAT
        ws.cell(row=r, column=5, value=f"=SUM(F{r}:N{r})").number_format = MONEY_FORMAT
        r += 1
    last = r - 1

    tot = r
    ws.cell(row=tot, column=4, value="TOTAL FEES AND EXPENSES").font = bold
    for c in range(5, 15):
        letter = ws.cell(row=tot, column=c).column_letter
        cell = ws.cell(row=tot, column=c, value=f"=SUM({letter}{first}:{letter}{last})" if last >= first else 0)
        cell.number_format, cell.font = MONEY_FORMAT, bold
    bp = log["totals"]["fee_bp"]
    rows = (
        (tot + 1, f"PRODUCTION FEE {fee_pct:g}%", f"=ROUND(E{tot}*{bp}/10000,2)"),
        (tot + 2, "GRAND TOTAL ALL FEES AND EXPENSES", f"=E{tot}+E{tot + 1}"),
        (tot + 3, "LESS ADVANCE RECEIVED", log["totals"]["advance_cents"] / 100),
        (tot + 4, "DIFFERENCE BETWEEN ADVANCE AND ACTUALS", f"=E{tot + 2}-E{tot + 3}"),
    )
    for row, label, value in rows:
        ws.cell(row=row, column=4, value=label).font = bold
        cell = ws.cell(row=row, column=5, value=value)
        cell.number_format, cell.font = MONEY_FORMAT, bold

    for letter, width in {"A": 5, "B": 10, "C": 26, "D": 40, "E": 13}.items():
        ws.column_dimensions[letter].width = width
    for c in range(6, 15):
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = 13
    ws.freeze_panes = "A9"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_env: Environment | None = None


def _environment() -> Environment:
    global _env
    if _env is None:
        _env = Environment(loader=FileSystemLoader(Path(__file__).parent / "templates"),
                           autoescape=select_autoescape(["html", "j2", "html.j2"], default=True),
                           trim_blocks=True, lstrip_blocks=True)
    return _env


def to_html(log: dict) -> str:
    cats = list(CATEGORIES)
    t = log["totals"]
    rows = [{"number": r.get("number"), "date": _date(r.get("spent_on")), "payee": r.get("payee") or "",
             "description": r.get("description") or "", "total": dollars(r["amount_cents"]),
             "cells": {k: dollars(r["amount_cents"]) if r["category"] == k else "" for k in cats}}
            for r in shown_rows(log)]
    sent_on = log.get("sent_on")
    return _environment().get_template("pc_log.html.j2").render(
        title=title(log), header=log["header"], cats=[(k, CATEGORIES[k]) for k in cats], rows=rows,
        by_category={k: dollars(v) if v else "" for k, v in t["by_category"].items()},
        total=dollars(t["total_cents"]), fee_pct=f"{t['fee_bp'] / 100:g}", fee=dollars(t["fee_cents"]),
        grand=dollars(t["grand_total_cents"]), advance=dollars(t["advance_cents"]),
        difference=dollars(t["difference_cents"]),
        sent=_date(sent_on) if sent_on else "")
