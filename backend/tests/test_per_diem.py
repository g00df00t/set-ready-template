"""The per diem sheet (services/per_diem.py).

    cd backend && venv/bin/python tests/test_per_diem.py

Checked against the producer's own sheet (docs/producer-documents-map.md §1.6, invented
rows): a traveller on the whole job gets 70+70+50+50+70+50+70 = 430; a local
who travels only for the second brand's leg gets 70+50+70 = 190; someone
booked for one shoot day gets 50. And it carries no rates, phones or emails.
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from services import per_diem  # noqa: E402
from services.job_dates import DatesError  # noqa: E402

FAILS = []


def check(name, ok, detail=""):
    print(f"  {'PASS' if ok else 'FAIL'}  {name}{'' if ok else '  -- ' + str(detail)}")
    if not ok:
        FAILS.append(name)


# Her job: Travel · Prep · Shoot A · Shoot A · Prep · Shoot B · Travel/Wrap
CREW = {"travel": [["2026-08-24", "2026-08-24"]],
        "prep": [["2026-08-25", "2026-08-25"], ["2026-08-28", "2026-08-28"]],
        "shoot": [["2026-08-26", "2026-08-27"], ["2026-08-29", "2026-08-29"]],
        "travel_out": [["2026-08-30", "2026-08-30"]]}
JOB = {"crew_dates": CREW, "per_diem_shoot": 50, "per_diem_other": 70,
       "day_labels": {"2026-08-26": "Brand A", "2026-08-29": "Brand B"}}
TRAVELLER = {"id": "t", "name": "Whole Job", "dates_override": {}, "phone": "555-0100",
             "email": "whole@example.com", "rate_shoot": 800}
LEG = {"id": "l", "name": "Local Except B", "away": [["2026-08-28", "2026-08-30"]], "dates_override": {},
       "per_diem_note": "local – except for brand B"}
ONE = {"id": "o", "name": "One Day", "dates_override": {
    "travel": [], "prep": [], "travel_out": [], "shoot": [["2026-08-27", "2026-08-27"]]}}
LOCAL = {"id": "x", "name": "Lives Here", "away": [], "dates_override": {}}

print("\nThe producer's worked check")
sheet = per_diem.build_sheet(JOB, [TRAVELLER, LEG, ONE, LOCAL])
by = {r["name"]: r for r in sheet["rows"]}
check("a traveller on the whole job gets 430", by["Whole Job"]["total"] == 430, by["Whole Job"]["cells"])
check("$50 on shoot days, $70 on every other day, prep included",
      by["Whole Job"]["cells"]["2026-08-26"] == 50 and by["Whole Job"]["cells"]["2026-08-25"] == 70
      and by["Whole Job"]["cells"]["2026-08-30"] == 70)
check("a local away for brand B's leg gets 190", by["Local Except B"]["total"] == 190, by["Local Except B"]["cells"])
check("a one-shoot-day booking gets 50 — the last away day is paid, unlike a hotel night",
      by["One Day"]["total"] == 50, by["One Day"]["cells"])
check("a local isn't on the sheet", "Lives Here" not in by)
check("the sheet total", sheet["total"] == 430 + 190 + 50)
check("a day's column total", sheet["totals"]["2026-08-27"] == 100, sheet["totals"])
check("every date in the run is a column", [d["date"] for d in sheet["dates"]][0] == "2026-08-24"
      and len(sheet["dates"]) == 7)
check("the header carries the day's label", sheet["dates"][2]["label"] == "Brand A")

print("\nDown days and check-ins")
s2 = per_diem.build_sheet({**JOB, "crew_dates": {**CREW, "off": [["2026-08-28", "2026-08-28"]], "prep": [["2026-08-25", "2026-08-25"]]}},
                          [{"id": "c", "name": "In Town", "dates_override": {"travel": [], "checkin": [["2026-08-24", "2026-08-24"]]}}])
cells = s2["rows"][0]["cells"]
check("a check-in day and a down day are $70 days", cells["2026-08-24"] == 70 and cells["2026-08-28"] == 70, cells)

print("\nA hotel night is a day away")
kept = per_diem.build_sheet(JOB, [{**LEG, "room_overrides": {"2026-08-26": "night"}}])["rows"][0]["cells"]
check("a night added on Rooming is paid, and so is the check-out morning after it",
      kept.get("2026-08-26") == 50 and kept.get("2026-08-27") == 50 and sum(kept.values()) == 190 + 100, kept)
gone = per_diem.build_sheet(JOB, [{**TRAVELLER, "room_overrides": {"2026-08-26": "none"}}])["rows"][0]
check("taking a hotel night away doesn't take their per diem away", gone["total"] == 430, gone["cells"])
check("a local with no nights still gets nothing",
      not per_diem.build_sheet(JOB, [LOCAL])["rows"])

print("\nHand edits and paid")
edited = per_diem.build_sheet(JOB, [{**TRAVELLER, "per_diem_overrides": {"2026-08-24": 0, "2026-08-25": 35}}])["rows"][0]
check("0 takes a day off, an amount replaces it",
      "2026-08-24" not in edited["cells"] and edited["cells"]["2026-08-25"] == 35 and edited["total"] == 430 - 70 - 35)
check("paid reads like her Initials column",
      per_diem.paid_text({"method": "venmo", "date": "2026-10-24"}) == "paid Venmo 10/24"
      and per_diem.paid_text({"method": "cash", "date": "2026-08-07", "note": "signed"}) == "paid cash 8/7 – signed")
check("not paid yet is blank", per_diem.paid_text(None) == "")

print("\nNo amounts set yet")
empty = per_diem.build_sheet({"crew_dates": CREW}, [TRAVELLER])
check("no amounts: no rows, and the sheet says why", empty["rows"] == [] and empty["missing_amounts"] is True)

print("\nWhat it must never carry")
row_keys = {k for r in sheet["rows"] for k in r}
flat = repr(sheet)
check("rows carry name, note, days, total and paid — no phone, email or rate",
      row_keys == {"hire_id", "name", "notes", "cells", "total", "paid", "paid_text"}
      and "555-0100" not in flat and "whole@example.com" not in flat and "800" not in flat, row_keys)

print("\nShape checks")
check("edits validate", per_diem.validate_overrides({"2026-08-24": "0", "2026-08-25": 35.5}) == {"2026-08-24": 0, "2026-08-25": 35.5})
for bad, fn, why in (({"2026-08-24": -5}, per_diem.validate_overrides, "a negative amount"),
                     ({"2026-08-24": 99999}, per_diem.validate_overrides, "an absurd amount"),
                     ({"method": "bitcoin"}, per_diem.validate_paid, "an unknown payment method")):
    try:
        fn(bad)
        check(f"refused: {why}", False)
    except DatesError:
        check(f"refused: {why}", True)
check("paid validates and trims", per_diem.validate_paid({"method": "zelle", "date": "2026-08-30", "note": "  sent  "})
      == {"method": "zelle", "date": "2026-08-30", "note": "sent"})


print("\nThe files")
import io  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from services.per_diem_documents import to_html, to_xlsx  # noqa: E402
job = {**JOB, "name": "Northwind Fall", "client_name": "Northwind"}
sheet = per_diem.build_sheet(job, [{**TRAVELLER, "per_diem_paid": {"method": "venmo", "date": "2026-08-30"}}, LEG, ONE])
values = [list(r) for r in load_workbook(io.BytesIO(to_xlsx(job, sheet))).active.iter_rows(values_only=True)]
check("the workbook is titled for the job, with the group and the rates",
      values[0][0] == "Per Diem - Northwind Fall" and values[1][:2] == ["Group:", "Northwind"]
      and "$70 non-shoot day · $50 shoot day" in values[2][1], values[:3])
heads = values[6]
check("each day's heading carries its type and label", "Shoot day Brand A" in heads and "Prep/Scout" in heads, heads)
row = next(r for r in values if r[0] == "Whole Job")
check("amounts are numbers with a live total, and Paid in the last column",
      row[2] == 70 and row[4] == 50 and str(row[9]).startswith("=SUM(") and row[10] == "paid Venmo 8/30", row)
check("a Totals row sums each column", any(r[0] == "Totals" and str(r[2]).startswith("=SUM(") for r in values))
html = to_html(job, sheet)
check("the PDF prints dollars and the totals", "$430" in html and "$670" in html and "paid Venmo 8/30" in html)
for name, text in (("Excel", repr(values)), ("PDF", html)):
    check(f"the {name} has no phone, email or rate",
          "555-0100" not in text and "whole@example.com" not in text and "$800" not in text and ">800<" not in text)

print("\nThe API")
from fastapi import HTTPException  # noqa: E402
from tests.fake_supabase import FakeClient  # noqa: E402
from routers import jobs as jobs_router, per_diem as pd_router  # noqa: E402
J = "00000000-0000-0000-0000-0000000000b1"
H = "00000000-0000-0000-0000-0000000000a1"
OTHER = "00000000-0000-0000-0000-0000000000b9"
db = {"jobs": [{"id": J, "owner_id": "o1", "name": "Northwind", "crew_dates": CREW, "per_diem_shoot": 50, "per_diem_other": 70},
               {"id": OTHER, "owner_id": "o2", "name": "Theirs", "crew_dates": CREW}],
      "job_hires": [{"id": H, "job_id": J, "name": "Whole Job", "dates_override": {}, "sort_order": 0}],
      "team_members": [], "crew_members": [], "roster": []}
client = FakeClient(db)
for mod in (jobs_router, pd_router):
    mod.get_client = lambda: client
got = pd_router.get_per_diem(J, owner_ids={"o1"})
check("the tab shows what's owed before anything is saved", got["rows"][0]["total"] == 430)
pd_router.update_hire_per_diem(J, H, pd_router.HirePerDiemUpdate(
    per_diem_overrides={"2026-08-24": 0}, per_diem_note=" flew in late ", per_diem_paid={"method": "cash", "date": "2026-08-30"}),
    owner_ids={"o1"})
row = pd_router.get_per_diem(J, owner_ids={"o1"})["rows"][0]
check("an edit, a note and a payment save and show",
      row["total"] == 360 and row["notes"] == "flew in late" and row["paid_text"] == "paid cash 8/30", row)
pd_router.update_hire_per_diem(J, H, pd_router.HirePerDiemUpdate(per_diem_paid=None), owner_ids={"o1"})
check("paid: null un-marks the payment (sent, not ignored)", db["job_hires"][0]["per_diem_paid"] is None)
try:
    pd_router.get_per_diem(OTHER, owner_ids={"o1"})
    check("someone else's job is refused", False)
except HTTPException as e:
    check("someone else's job is refused", e.status_code == 404)
resp = pd_router.download_per_diem_xlsx(J, owner_ids={"o1"})
check("the Excel download names the job", "Per%20Diem%20-%20Northwind" in resp.headers["content-disposition"])
db["jobs"][0]["name"] = "Rates test — second  job"
cd = pd_router.download_per_diem_xlsx(J, owner_ids={"o1"}).headers["content-disposition"]
check("a dash in the job name stays a dash, no double spaces", "Rates%20test%20-%20second%20job.xlsx" in cd, cd)

print()
print(f"{len(FAILS)} FAILED: {', '.join(FAILS)}" if FAILS else "ALL PASS")
sys.exit(1 if FAILS else 0)
