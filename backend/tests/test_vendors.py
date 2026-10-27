"""Vendor bookings, usage dates and the job's contact sheet
(services/vendors.py, services/contact_sheet.py, routers/vendors.py,
routers/job_info.py).

    cd backend && venv/bin/python tests/test_vendors.py

The producer's Job Info Sheet (docs/producer-documents-map.md §1.9): the hotel,
each day's location, permits and the motorhome each carry COI requested ·
Signed agreement · Deposit · Final bill. Patrick, 2026-10-02: link to what the
job already holds (Rooming's stays, the days' locations), a date per step with
amounts on the payments, and the contact sheet as a view. Everything invented.
"""

import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from fastapi import HTTPException  # noqa: E402

from services import contact_sheet, vendors as svc  # noqa: E402
from services.job_dates import DatesError, talent_dates  # noqa: E402
from tests.fake_supabase import FakeClient  # noqa: E402
from routers import job_info as info_router, vendors as vendors_router  # noqa: E402
import routers.jobs as jobs_router  # noqa: E402  (the _own_job both routers use)

FAILS = []


def check(name, ok, detail=""):
    print(f"  {'PASS' if ok else 'FAIL'}  {name}{'' if ok else '  -- ' + str(detail)}")
    if not ok:
        FAILS.append(name)


def refused(name, fn, *args):
    try:
        fn(*args)
        check(name, False, "not refused")
    except DatesError:
        check(name, True)


J = "00000000-0000-0000-0000-0000000000e1"
STAY = {"id": "s1", "name": "The Oxford", "address": "1 Main St", "check_in": "2026-10-12", "check_out": "2026-10-16", "rooms": []}
CREW = {"travel": [["2026-10-12", "2026-10-12"]], "shoot": [["2026-10-13", "2026-10-14"]],
        "travel_out": [["2026-10-15", "2026-10-15"]]}
HOTELS, DAYS = {"s1"}, {"d1", "d2"}

print("\nA booking's shape")
v = svc.validate_vendor({"kind": "permit", "name": "  City  film permit ", "fee": "450", "contact_email": "permits@example.com",
                         "website": "film.example.gov", "notes": "Apply 10 days out\n\n  bring COI ",
                         "steps": {"coi": {"date": "2026-10-01"}, "deposit": {"date": "2026-10-02", "amount": "100"},
                                   "agreement": None}}, HOTELS, DAYS)
check("text is tidied, the fee is a number", v["name"] == "City film permit" and v["fee"] == 450.0)
check("a website without a scheme becomes a link", v["website"] == "https://film.example.gov")
check("notes keep their lines", v["notes"] == "Apply 10 days out\n\nbring COI")
check("a step is done when it has a date; payments carry an amount",
      v["steps"] == {"coi": {"date": "2026-10-01"}, "deposit": {"date": "2026-10-02", "amount": 100.0}}, v["steps"])
check("it gets an id", bool(v["id"]))
h = svc.validate_vendor({"kind": "hotel", "hotel_id": "s1", "name": "ignored? no — kept", "day_id": "d1"}, HOTELS, DAYS)
check("a hotel booking links to a stay, never a day", h["hotel_id"] == "s1" and h["day_id"] is None)
loc = svc.validate_vendor({"kind": "location", "day_id": "d2"}, HOTELS, DAYS)
check("a location booking links to a day, needs no name", loc["day_id"] == "d2" and loc["name"] is None)
cat = svc.validate_vendor({"kind": "caterer", "name": "Taco Truck", "steps": {"coi": {"date": "2026-10-01"}}}, HOTELS, DAYS)
check("a caterer has no checklist (as on her sheet)", cat["steps"] == {})
refused("an unknown kind", svc.validate_vendor, {"kind": "spaceship", "name": "x"}, HOTELS, DAYS)
refused("a hotel booking with no stay", svc.validate_vendor, {"kind": "hotel"}, HOTELS, DAYS)
refused("a hotel booking for someone else's stay", svc.validate_vendor, {"kind": "hotel", "hotel_id": "zz"}, HOTELS, DAYS)
refused("a location booking with no day", svc.validate_vendor, {"kind": "location"}, HOTELS, DAYS)
refused("a permit with no name", svc.validate_vendor, {"kind": "permit"}, HOTELS, DAYS)
refused("a negative fee", svc.validate_vendor, {"kind": "permit", "name": "x", "fee": -5}, HOTELS, DAYS)
refused("a fee that isn't a number", svc.validate_vendor, {"kind": "permit", "name": "x", "fee": "lots"}, HOTELS, DAYS)
refused("an unknown checklist step", svc.validate_vendor, {"kind": "permit", "name": "x", "steps": {"vibes": {"date": "2026-10-01"}}}, HOTELS, DAYS)
refused("a bad step date", svc.validate_vendor, {"kind": "permit", "name": "x", "steps": {"coi": {"date": "soon"}}}, HOTELS, DAYS)

print("\nLinked, not copied")
job = {"hotels": [STAY]}
days = [{"id": "d1", "name": "Day 1", "shoot_date": "2026-10-13"}, {"id": "d2", "name": "Day 2", "shoot_date": "2026-10-14"}]
locations = {"d2": {"project_id": "d2", "location_name": "Studio 4", "address": "200 Water St"}}
r = {x["kind"]: x for x in svc.resolve(job, [h, loc, v, cat], days, locations)}
check("the hotel booking reads the stay's name and dates", r["hotel"]["title"] == "The Oxford"
      and r["hotel"]["subtitle"] == "Oct 12 – Oct 16 · 1 Main St", r["hotel"])
check("...and a rename on Rooming shows here", svc.resolve({"hotels": [{**STAY, "name": "The Ace"}]}, [h], days, locations)[0]["title"] == "The Ace")
check("the location booking reads the day's location", r["location"]["title"] == "Studio 4"
      and r["location"]["subtitle"] == "Oct 14 · 200 Water St", r["location"])
check("a booking whose stay was removed is marked broken", svc.resolve({"hotels": []}, [h], days, locations)[0]["broken"])
check("missing steps are listed in order", r["permit"]["missing"] == ["agreement", "final_bill"])
check("a caterer is never 'missing' anything", r["caterer"]["missing"] == [])
check("bookings sort by kind: hotel, location, permit, caterer",
      [x["kind"] for x in svc.resolve(job, [cat, v, loc, h], days, locations)] == ["hotel", "location", "permit", "caterer"])
opened = svc.open_items(svc.resolve(job, [h, loc, v, cat], days, locations))
check("open items: what's still to do, by booking",
      [(o["title"], o["missing"]) for o in opened][2] == ("City film permit", ["Signed agreement", "Final bill (balance)"]), opened)
check("...a broken booking isn't nagged about", svc.open_items(svc.resolve({"hotels": []}, [h], days, locations)) == [])
ub = svc.unbooked(job, [], days, locations)
check("unbooked: the stay and the day with a location are offered",
      [x["hotel_id"] for x in ub["hotels"]] == ["s1"] and [x["day_id"] for x in ub["days"]] == ["d2"], ub)
ub = svc.unbooked(job, [h, loc], days, locations)
check("...and stop being offered once booked", ub == {"hotels": [], "days": []})

print("\nUsage dates")
check("four dates, blanks dropped", svc.validate_usage({"first_use": "2027-03-01", "launch": "", "web_expiry": None})
      == {"first_use": "2027-03-01"})
refused("an unknown usage date", svc.validate_usage, {"birthday": "2027-01-01"})
refused("a bad usage date", svc.validate_usage, {"launch": "spring"})

print("\nThe contact sheet")
HIRES = [
    {"id": "c1", "name": "Sam Ortiz", "role": "Gaffer", "phone": "555-0101", "email": "sam@example.com",
     "dates_override": {}, "sort_order": 0},
    {"id": "c2", "name": "Held Crew", "role": "PA", "status": "hold", "hold_rank": 2, "dates_override": {}, "sort_order": 1},
    {"id": "c3", "name": "Gone", "role": "Grip", "status": "released", "dates_override": {}, "sort_order": 2},
    {"id": "t1", "kind": "talent", "name": "Dee Park", "phone": "555-0102", "email": "dee@example.com", "roster_id": "r1",
     "segment": "Women", "dates_override": talent_dates(["2026-10-14"]), "sort_order": 3},
]
ROSTER = {"r1": {"id": "r1", "agency": "Wilhelmina", "agent_name": "Ana Agent", "agent_phone": "555-0300",
                 "agent_email": "ana@example.com"}}
CONTACTS = [
    {"project_id": "d1", "organization": "Northwind", "name": "Casey Client", "role": "Art Director", "email": "casey@example.com"},
    {"project_id": "d2", "organization": "Northwind", "name": "Casey Client", "role": "", "phone": "555-0400", "email": "CASEY@example.com"},
    {"project_id": "d2", "organization": "Agency X", "name": "No Email", "role": "Producer"},
]
cjob = {"crew_dates": CREW}
vend = svc.resolve(job, [{**v, "contact_name": "Pat Permit", "contact_phone": "555-0500"}, h], days, locations)
sheet = contact_sheet.build(cjob, HIRES, ROSTER, CONTACTS, vend)
check("crew with their travel-to-travel dates", sheet["crew"][0]["first"] == "2026-10-12" and sheet["crew"][0]["last"] == "2026-10-15")
check("a hold shows as a hold; a release is left off",
      [(c["name"], c["status"]) for c in sheet["crew"]] == [("Sam Ortiz", ""), ("Held Crew", "2nd hold")])
t = sheet["talent"][0]
check("talent with their own dates, agency and agent", (t["first"], t["last"], t["agency"], t["agent"], t["agent_phone"])
      == ("2026-10-13", "2026-10-15", "Wilhelmina", "Ana Agent", "555-0300"), t)
check("a client on two days is one row, filled from both",
      len([c for c in sheet["clients"] if c["name"] == "Casey Client"]) == 1
      and next(c for c in sheet["clients"] if c["name"] == "Casey Client")["phone"] == "555-0400")
check("a contact with no email is kept, by name", any(c["name"] == "No Email" for c in sheet["clients"]))
check("vendors with a contact are listed; one without isn't",
      [(x["title"], x["name"]) for x in sheet["vendors"]] == [("City film permit", "Pat Permit")], sheet["vendors"])
from openpyxl import load_workbook  # noqa: E402
wb = load_workbook(io.BytesIO(contact_sheet.to_xlsx({"name": "Fall"}, sheet, [("Launch date:", "2027-03-01")])))
vals = [c for row in wb.active.iter_rows(values_only=True) for c in row if c]
check("the Excel: title, usage dates, all four blocks",
      vals[0] == "Job Info - Fall" and "2027-03-01" in vals and all(b in vals for b in ("CREW", "TALENT", "CLIENT & AGENCY", "VENDORS")))
check("...with the travel-to-travel dates", "2026-10-12 to 2026-10-15" in vals)

print("\nThrough the API")
OTHER = "00000000-0000-0000-0000-0000000000e9"
db = {"jobs": [{"id": J, "owner_id": "o1", "name": "Fall", "crew_dates": CREW, "hotels": [STAY]},
               {"id": OTHER, "owner_id": "o2", "name": "Theirs", "crew_dates": {}, "hotels": []}],
      "projects": [{"id": "d1", "owner_id": "o1", "job_id": J, "name": "Day 1", "shoot_date": "2026-10-13"},
                   {"id": "d2", "owner_id": "o1", "job_id": J, "name": "Day 2", "shoot_date": "2026-10-14"}],
      "project_locations": [locations["d2"]], "project_contacts": [dict(c) for c in CONTACTS],
      "job_hires": [{**x, "job_id": J} for x in HIRES], "roster": [{**ROSTER["r1"], "owner_id": "o1"}]}
client = FakeClient(db)
vendors_router.get_client = info_router.get_client = jobs_router.get_client = lambda: client
B = vendors_router.VendorBody
out = vendors_router.get_vendors(J, owner_ids={"o1"})
check("a job before the migration reads as no bookings", out["vendors"] == [] and len(out["unbooked"]["hotels"]) == 1)
out = vendors_router.add_vendor(J, B(vendor={"kind": "hotel", "hotel_id": "s1", "fee": 1800}), owner_ids={"o1"})
hid = out["vendors"][0]["id"]
check("adding the hotel booking", out["vendors"][0]["title"] == "The Oxford" and out["unbooked"]["hotels"] == [])
vendors_router.add_vendor(J, B(vendor={"kind": "motorhome", "name": "Star Coach"}), owner_ids={"o1"})
out = vendors_router.update_vendor(J, hid, B(vendor={"kind": "hotel", "hotel_id": "s1", "fee": 1800,
                                                     "steps": {"coi": {"date": "2026-10-01"}}}), owner_ids={"o1"})
check("ticking a step saves its date, the other booking untouched",
      out["vendors"][0]["steps"]["coi"]["date"] == "2026-10-01" and len(out["vendors"]) == 2)
check("open items count both bookings' remaining steps", len(out["open_items"]) == 2 and
      out["open_items"][0]["missing"] == ["Signed agreement", "Deposit", "Final bill (balance)"], out["open_items"])
try:
    vendors_router.add_vendor(J, B(vendor={"kind": "location", "day_id": "not-ours"}), owner_ids={"o1"})
    check("a location on another job's day is a 422", False)
except HTTPException as e:
    check("a location on another job's day is a 422", e.status_code == 422)
try:
    vendors_router.update_vendor(J, "nope", B(vendor={"kind": "other", "name": "x"}), owner_ids={"o1"})
    check("editing a booking that isn't there is a 404", False)
except HTTPException as e:
    check("editing a booking that isn't there is a 404", e.status_code == 404)
out = vendors_router.delete_vendor(J, hid, owner_ids={"o1"})
check("removing a booking leaves the rest", [x["title"] for x in out["vendors"]] == ["Star Coach"])
try:
    vendors_router.get_vendors(OTHER, owner_ids={"o1"})
    check("someone else's job is refused", False)
except HTTPException as e:
    check("someone else's job is refused", e.status_code == 404)
saved = info_router.update_usage_dates(J, info_router.UsageBody(usage_dates={"launch": "2027-03-01"}), owner_ids={"o1"})
check("usage dates save", saved == {"launch": "2027-03-01"})
info = info_router.get_job_info(J, owner_ids={"o1"})
check("the Job info tab: usage dates and the contact sheet",
      info["usage_dates"] == {"launch": "2027-03-01"} and [c["name"] for c in info["crew"]] == ["Sam Ortiz", "Held Crew"]
      and info["talent"][0]["agency"] == "Wilhelmina")
resp = info_router.download_job_info(J, owner_ids={"o1"})
check("the Excel downloads, named for the job", "Job%20Info%20-%20Fall.xlsx" in resp.headers["content-disposition"])

print()
print(f"{len(FAILS)} FAILED: {', '.join(FAILS)}" if FAILS else "ALL PASS")
sys.exit(1 if FAILS else 0)
