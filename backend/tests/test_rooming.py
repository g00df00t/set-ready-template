"""The rooming list (services/rooming.py).

    cd backend && venv/bin/python tests/test_rooming.py

The rules from docs/producer-documents-map.md §2.2–§2.3 and the producer's answers of
2026-10-01: a night is an away date that isn't the last of its run; a single
away day has no room; a hotel move mid-job splits the nights between stays;
hand edits add a night, take one away, or mark a late check-out. And because
the list goes to the hotel, no row ever carries a phone, an email, a rate or
per diem.
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from services import rooming  # noqa: E402
from services.job_dates import DatesError  # noqa: E402

FAILS = []


def check(name, ok, detail=""):
    print(f"  {'PASS' if ok else 'FAIL'}  {name}{'' if ok else '  -- ' + str(detail)}")
    if not ok:
        FAILS.append(name)


CREW = {"travel": [["2026-10-19", "2026-10-19"]], "shoot": [["2026-10-20", "2026-10-20"], ["2026-10-22", "2026-10-22"]],
        "off": [["2026-10-21", "2026-10-21"]], "travel_out": [["2026-10-23", "2026-10-23"]]}

print("\nNights from who's away")
ns, late = rooming.nights(CREW, {}, None, {})
check("a traveller sleeps every night but the last, which is check-out",
      ns == ["2026-10-19", "2026-10-20", "2026-10-21", "2026-10-22"] and not late, ns)
check("a local has no nights", rooming.nights(CREW, {}, [], {})[0] == [])
ns, _ = rooming.nights(CREW, {}, [["2026-10-21", "2026-10-23"]], {})
check("away only for a leg: nights inside it only", ns == ["2026-10-21", "2026-10-22"], ns)
ns, _ = rooming.nights(CREW, {}, [["2026-10-22", "2026-10-22"]], {})
check("a single away day (in and out) has no room", ns == [], ns)
ns, late = rooming.nights(CREW, {}, None, {"2026-10-23": "night", "2026-10-21": "none", "2026-10-24": "late"})
check("hand edits: add a night, take one away, mark a late check-out",
      ns == ["2026-10-19", "2026-10-20", "2026-10-22", "2026-10-23"] and late == {"2026-10-24"}, (ns, late))
check("a person with their own dates gets nights from those",
      rooming.nights(CREW, {"travel": [], "checkin": [["2026-10-19", "2026-10-19"]]}, None, {})[0][0] == "2026-10-19")

print("\nOne sheet per hotel stay")
HOTEL_A = {"id": "a", "name": "Hotel A", "check_in": "2026-10-19", "check_out": "2026-10-21"}
HOTEL_B = {"id": "b", "name": "Hotel B", "check_in": "2026-10-21", "check_out": "2026-10-23"}
job = {"crew_dates": CREW, "hotels": [HOTEL_A, HOTEL_B], "day_labels": {"2026-10-20": "Brand A"}}
SAM = {"id": "h1", "name": "Sam Ortiz", "email": "sam@example.com", "phone": "555-0100", "rate_shoot": 800,
       "room_notes": "late arrival", "reservations": {"a": "A-123", "b": "B-9"}, "dates_override": {}}
LOCAL = {"id": "h2", "name": "Lee Local", "away": [], "dates_override": {}}
out = rooming.build_sheets(job, [SAM, LOCAL])
a, b = out["sheets"]
check("a hotel move splits the nights between the two stays",
      a["rows"][0]["cells"] == {"2026-10-19": "1", "2026-10-20": "1", "2026-10-21": "c/o"}
      and b["rows"][0]["cells"] == {"2026-10-21": "1", "2026-10-22": "1", "2026-10-23": "c/o"}, (a["rows"], b["rows"]))
check("each stay shows its own reservation number", a["rows"][0]["reservation"] == "A-123" and b["rows"][0]["reservation"] == "B-9")
check("locals aren't on the list", all(len(s["rows"]) == 1 for s in out["sheets"]))
check("totals count rooms a night", a["totals"] == {"2026-10-19": 1, "2026-10-20": 1, "2026-10-21": 0} and a["total"] == 2, a["totals"])
check("the header carries the crew's day types and the day's label",
      a["dates"][1] == {"date": "2026-10-20", "types": ["shoot"], "label": "Brand A"}, a["dates"][1])
check("every night has a hotel", out["unplaced"] == [])

row_keys = {k for s in out["sheets"] for r in s["rows"] for k in r}
flat = repr(out)
check("⚠️ no row carries a phone, an email, a rate or per diem — the list goes to the hotel",
      row_keys == {"kind", "hire_id", "room_id", "name", "notes", "reservation", "cells", "total"}
      and "555-0100" not in flat and "sam@example.com" not in flat and "800" not in flat, row_keys)

out = rooming.build_sheets({**job, "hotels": [HOTEL_A]}, [SAM])
check("nights no stay covers are reported, not dropped",
      [u["date"] for u in out["unplaced"]] == ["2026-10-21", "2026-10-22"], out["unplaced"])
out = rooming.build_sheets({**job, "hotels": []}, [SAM])
check("no hotels yet: one sheet over every night, check-out included",
      len(out["sheets"]) == 1 and out["sheets"][0]["hotel"] is None
      and out["sheets"][0]["rows"][0]["cells"].get("2026-10-23") == "c/o", out["sheets"])
late = rooming.build_sheets(job, [{**SAM, "room_overrides": {"2026-10-23": "late"}}])["sheets"][1]
check("a late check-out prints as such", late["rows"][0]["cells"]["2026-10-23"] == "late c/o")

print("\nShape checks")
hotels = rooming.validate_hotels([{"name": "  Hotel   B ", "check_in": "2026-10-21", "check_out": "2026-10-23"},
                                  {"id": "keep", "name": "Hotel A", "check_in": "2026-10-19", "check_out": "2026-10-21"}])
check("hotels sort by date, keep their ids, get one when new, trim text",
      hotels[0]["id"] == "keep" and hotels[1]["id"] and hotels[1]["name"] == "Hotel B", hotels)
for bad, why in (([{"check_in": "2026-10-21", "check_out": "2026-10-21"}], "check-out after check-in"),
                 ({"2026-10-21": "maybe"}, "room edit values are fixed")):
    try:
        (rooming.validate_hotels if isinstance(bad, list) else rooming.validate_overrides)(bad)
        check(f"refused: {why}", False)
    except DatesError:
        check(f"refused: {why}", True)
check("away: None stays None, [] stays local", rooming.validate_away(None) is None and rooming.validate_away([]) == [])


print("\nExtra rooms (the producer's merch room)")
MERCH = {"id": "m", "label": "Merch room", "notes": "racks", "reservation": "A-77",
         "nights": ["2026-10-19", "2026-10-20"], "late": ["2026-10-21"]}
out = rooming.build_sheets({**job, "hotels": [{**HOTEL_A, "rooms": [MERCH]}, HOTEL_B]}, [SAM])
merch = [r for r in out["sheets"][0]["rows"] if r["kind"] == "room"][0]
check("a room's nights, then its (late) check-out",
      merch["cells"] == {"2026-10-19": "1", "2026-10-20": "1", "2026-10-21": "late c/o"} and merch["total"] == 2, merch)
check("it prints its own label, note and reservation", (merch["name"], merch["notes"], merch["reservation"]) == ("Merch room", "racks", "A-77"))
check("rooms a night count it; breakfasts don't",
      out["sheets"][0]["totals"]["2026-10-19"] == 2 and out["sheets"][0]["breakfasts"]["2026-10-19"] == 1)
rooms = rooming.validate_rooms([{"label": "  Merch   room ", "nights": ["2026-10-20", "2026-10-19", "2026-10-19"]}],
                               "2026-10-19", "2026-10-21")
check("rooms validate: trimmed, nights deduped and sorted, an id given",
      rooms[0]["label"] == "Merch room" and rooms[0]["nights"] == ["2026-10-19", "2026-10-20"] and rooms[0]["id"])
for bad, why in (([{"label": "", "nights": []}], "a room needs a name"),
                 ([{"label": "Merch", "nights": ["2026-10-21"]}], "a night on check-out day isn't a night"),
                 ([{"label": "Merch", "late": ["2026-10-19"]}], "late check-out can't be on check-in day")):
    try:
        rooming.validate_rooms(bad, "2026-10-19", "2026-10-21")
        check(f"refused: {why}", False)
    except DatesError:
        check(f"refused: {why}", True)
from services.rooming_documents import to_html as _html, to_xlsx as _xlsx  # noqa: E402
import io as _io  # noqa: E402
from openpyxl import load_workbook as _load  # noqa: E402
withroom = {**job, "name": "F", "hotels": [{**HOTEL_A, "rooms": [MERCH]}, HOTEL_B]}
s0 = rooming.build_sheets(withroom, [SAM])["sheets"][0]
xrows = [r for r in _load(_io.BytesIO(_xlsx(withroom, s0))).active.iter_rows(values_only=True)]
check("the Excel adds a breakfast count when there's an extra room",
      any(r[0] == "Merch room" for r in xrows) and any(r[0] == "Hotel Breakfast Count" and r[3] == 1 for r in xrows))
check("...and the PDF", "Hotel Breakfast Count" in _html(withroom, s0) and "Merch room" in _html(withroom, s0))
check("no extra room, no breakfast row (it would just repeat the totals)",
      "Hotel Breakfast Count" not in _html(withroom, rooming.build_sheets(withroom, [SAM])["sheets"][1]))
check("rooms survive a hotels save",
      rooming.validate_hotels([{**HOTEL_A, "rooms": [MERCH]}])[0]["rooms"][0]["id"] == "m")

print("\nThe files for the hotel")
import io  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from services.rooming_documents import to_html, to_xlsx  # noqa: E402
job_full = {**job, "name": "Fall Lookbook", "client_name": "Riverbend Outdoor",
            "hotels": [{**HOTEL_A, "contact": "Front desk"}, HOTEL_B]}
sheet = rooming.build_sheets(job_full, [SAM, LOCAL])["sheets"][0]
ws = load_workbook(io.BytesIO(to_xlsx(job_full, sheet))).active
values = [[c for c in row] for row in ws.iter_rows(values_only=True)]
flat_x = repr(values)
check("the workbook opens, titled for the hotel", values[0][0] == "Hotel Rooming List - Hotel A", values[0])
check("Group and Contact lines like the producer's", values[1][:2] == ["Group:", "Riverbend Outdoor"] and values[2][:2] == ["Contact:", "Front desk"], values[:3])
row = next(r for r in values if r[0] == "Sam Ortiz")
check("a night is 1 and check-out is c/o, with a live TOTAL",
      row[3] == 1 and row[4] == 1 and row[5] == "c/o" and str(row[6]).startswith("=SUM("), row)
check("a TOTALS row sums each column", any(r[0] == "TOTALS" and str(r[3]).startswith("=SUM(") for r in values))
html = to_html(job_full, sheet)
for name, text in (("Excel", flat_x), ("PDF", html)):
    check(f"⚠️ the {name} has no phone, email or rate",
          "555-0100" not in text and "sam@example.com" not in text and "800" not in text)
check("the PDF escapes what a producer typed", "&lt;b&gt;" in to_html(job_full, {**sheet, "rows": [{**sheet["rows"][0], "notes": "<b>x</b>"}]}))

print("\nThe API")
from fastapi import HTTPException  # noqa: E402
from tests.fake_supabase import FakeClient  # noqa: E402
from routers import jobs as jobs_router, rooming as rooming_router  # noqa: E402
J = "00000000-0000-0000-0000-0000000000b1"
H = "00000000-0000-0000-0000-0000000000a1"
OTHER = "00000000-0000-0000-0000-0000000000b9"
db = {"jobs": [{"id": J, "owner_id": "o1", "name": "Fall", "crew_dates": CREW, "hotels": []},
               {"id": OTHER, "owner_id": "o2", "name": "Theirs", "crew_dates": {}, "hotels": []}],
      "job_hires": [{"id": H, "job_id": J, "name": "Sam Ortiz", "phone": "555-0100", "dates_override": {}, "sort_order": 0}],
      "team_members": [], "crew_members": [], "roster": []}
client = FakeClient(db)
for mod in (jobs_router, rooming_router):
    mod.get_client = lambda: client
saved = rooming_router.update_hotels(J, rooming_router.HotelsUpdate(hotels=[
    {"name": "Hotel A", "check_in": "2026-10-19", "check_out": "2026-10-23"}]), owner_ids={"o1"})
check("hotels save with an id", len(saved) == 1 and saved[0]["id"])
rooming_router.update_hire_rooming(J, H, rooming_router.HireRoomingUpdate(
    room_notes=" late  arrival ", reservations={saved[0]["id"]: "A-1"}, room_overrides={"2026-10-23": "late"}),
    owner_ids={"o1"})
got = rooming_router.get_rooming(J, owner_ids={"o1"})
row = got["sheets"][0]["rows"][0]
check("the tab shows the note, reservation and late check-out",
      row["notes"] == "late arrival" and row["reservation"] == "A-1" and row["cells"]["2026-10-23"] == "late c/o", row)
check("the people list carries their dates for the Some days picker",
      got["people"][0]["dates"][0] == "2026-10-19" and got["people"][0]["away"] is None)
rooming_router.update_hire_rooming(J, H, rooming_router.HireRoomingUpdate(away=[]), owner_ids={"o1"})
check("marking someone local takes them off the list",
      rooming_router.get_rooming(J, owner_ids={"o1"})["sheets"][0]["rows"] == [])
rooming_router.update_hire_rooming(J, H, rooming_router.HireRoomingUpdate(away=None), owner_ids={"o1"})
check("away: null puts them back on every date (sent, not ignored)", db["job_hires"][0]["away"] is None
      and rooming_router.get_rooming(J, owner_ids={"o1"})["sheets"][0]["rows"] != [])
for fn, args in ((rooming_router.get_rooming, (OTHER,)),
                 (rooming_router.update_hotels, (OTHER, rooming_router.HotelsUpdate(hotels=[])))):
    try:
        fn(*args, owner_ids={"o1"})
        check(f"someone else's job: {fn.__name__} refused", False)
    except HTTPException as e:
        check(f"someone else's job: {fn.__name__} refused", e.status_code == 404)
resp = rooming_router.download_rooming_xlsx(J, saved[0]["id"], owner_ids={"o1"})
check("the Excel download is a spreadsheet with a filename",
      resp.media_type.endswith("spreadsheetml.sheet") and "Hotel%20Rooming%20List" in resp.headers["content-disposition"])


hid = saved[0]["id"]
rooming_router.update_rooms(J, hid, rooming_router.RoomsUpdate(rooms=[{"label": "Merch room", "nights": ["2026-10-19"]}]),
                            owner_ids={"o1"})
sheet = rooming_router.get_rooming(J, owner_ids={"o1"})["sheets"][0]
check("a room added through the API shows on its hotel's list", any(r["kind"] == "room" and r["name"] == "Merch room" for r in sheet["rows"]))
check("saving rooms keeps the stay's other fields", db["jobs"][0]["hotels"][0]["name"] == "Hotel A")
try:
    rooming_router.update_rooms(J, "nope", rooming_router.RoomsUpdate(rooms=[]), owner_ids={"o1"})
    check("rooms for a hotel that isn't on the job are refused", False)
except HTTPException as e:
    check("rooms for a hotel that isn't on the job are refused", e.status_code == 404)

print()
print(f"{len(FAILS)} FAILED: {', '.join(FAILS)}" if FAILS else "ALL PASS")
sys.exit(1 if FAILS else 0)
