"""The producer's PC log: generated lines, hand changes, Mark as sent, the files
(services/pc_log.py, services/pc_log_documents.py, routers/pc_log.py).

    cd backend && venv/bin/python tests/test_pc_log.py

Decided 2026-10-02 (memory: reconciliation build): generated lines — each
booked person's rates × days, the per diem total, every vendor payment — are
live until the log is marked sent; then they're flagged, never applied. A day
is paid once: shoot, prep, WRAP AT THE PREP RATE, travel, down day at ½ shoot,
check-in unpaid. Talent add their agency fee. Every person and amount invented.
"""

import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from fastapi import HTTPException  # noqa: E402

from services import pc_log  # noqa: E402
from services.pc_log_documents import to_html, to_xlsx  # noqa: E402
from tests.fake_supabase import FakeClient  # noqa: E402
from routers import pc_log as pc_router, budget as budget_router, vendors as vendors_router  # noqa: E402
from models.budget import ExpenseCreate, InvoiceCreate  # noqa: E402
import routers.jobs as jobs_router  # noqa: E402

FAILS = []


def check(name, ok, detail=""):
    print(f"  {'PASS' if ok else 'FAIL'}  {name}{'' if ok else '  -- ' + str(detail)}")
    if not ok:
        FAILS.append(name)


J = "00000000-0000-0000-0000-0000000000c1"
CREW = {"travel": [["2026-10-11", "2026-10-11"]], "prep": [["2026-10-12", "2026-10-12"]],
        "shoot": [["2026-10-13", "2026-10-14"]], "off": [["2026-10-15", "2026-10-15"]],
        "wrap": [["2026-10-16", "2026-10-16"]], "travel_out": [["2026-10-17", "2026-10-17"]]}
RATES = {"rate_shoot": 1200, "rate_travel": 600, "rate_prep": 600}


def hire(i, **kw):
    return {"id": f"h{i}", "job_id": J, "name": f"Person {i}", "role": None, "department": None, "kind": "crew",
            "status": "confirmed", "dates_override": {}, "away": None, "per_diem_overrides": {}, **kw}


A = hire(1, name="Avery Sample", role="Photographer", department="Camera Department", **RATES)
# Their travel-in is on a shoot day: that date is paid once, at the shoot rate.
B = hire(2, name="Blake Sample", role="Sound mixer", department="Sound Department",
         dates_override={"travel": [["2026-10-13", "2026-10-13"]]}, **RATES)
C = hire(3, name="Casey Sample", role="HMU", department="Hair & Makeup", rate_shoot=900, rate_prep=450)
T = hire(4, name="Dana Sample", kind="talent", rate_shoot=1000, agency_fee_pct="20.00",
         dates_override={"shoot": [["2026-10-13", "2026-10-13"]]})
HOLD = hire(5, name="Ellis Sample", status="hold", **RATES)
GONE = hire(6, name="Finley Sample", status="released", **RATES)
JOB = {"id": J, "owner_id": "o1", "name": "Riverbend Fall", "client_name": "Example Co", "crew_dates": CREW,
       "per_diem_shoot": 50, "per_diem_other": 70, "production_fee_bp": 800, "pc_log_camera": "photography",
       "pc_log_generated": {},
       "hotels": [{"id": "s1", "name": "Example Inn", "address": "1 Main St", "check_in": "2026-10-11",
                   "check_out": "2026-10-17", "rooms": []}],
       "vendors": [
           {"id": "v1", "kind": "hotel", "hotel_id": "s1", "name": None,
            "steps": {"deposit": {"date": "2026-10-01", "amount": 500}, "final_bill": {"date": "2026-10-20", "amount": 1300.5}}},
           {"id": "v2", "kind": "location", "day_id": "d1", "name": None,
            "steps": {"final_bill": {"date": "2026-10-13", "amount": 2400}}},
           {"id": "v3", "kind": "caterer", "name": "Example Catering", "steps": {}},
           {"id": "v4", "kind": "permit", "name": "City permit", "steps": {"deposit": {"date": "2026-10-02", "amount": 0}}},
           {"id": "v5", "kind": "hotel", "hotel_id": "gone", "name": None,
            "steps": {"final_bill": {"date": "2026-10-20", "amount": 999}}}]}

print("\nWhich column a person's fee goes in")
for h, want in ((A, "photography"), (B, "video"), (C, "stylist_hmu"), (T, "talent"),
                ({"department": "Grip & Electric"}, "photography"), ({"department": "Art Department"}, "art_props"),
                ({"department": "Wardrobe/Costume"}, "stylist_hmu"), ({"role": "Location scout"}, "location_studio"),
                ({"department": "Production", "role": "PA"}, "production_misc"), ({}, "production_misc"),
                ({"role": "Director of Photography"}, "video")):
    check(f"{h.get('department') or h.get('role') or h.get('kind') or 'nothing'} -> {want}",
          pc_log.category_for_hire(h) == want, pc_log.category_for_hire(h))
check("the camera team follows the job's choice", pc_log.category_for_hire({"department": "Camera Department"}, "video") == "video")

print("\nA person's fee: their days at their rates")
fa = pc_log.hire_fee(JOB, A)
# 2 shoot × 1,200 + prep 600 + wrap (prep rate) 600 + 2 travel × 600 + down day ½ × 1,200
check("every type at its rate: 5,400.00", fa["amount_cents"] == 540_000, fa)
check("the line says how", fa["detail"] == "2 shoot x $1,200.00 + 1 prep x $600.00 + 1 wrap x $600.00 + 2 travel x $600.00 + 1 down x $600.00", fa["detail"])
check("travel in and out read as one", fa["detail"].count("travel") == 1, fa["detail"])
check("dated their last day", fa["last_date"] == "2026-10-17")
fb = pc_log.hire_fee(JOB, B)
check("a date with two types is paid once, at the shoot rate: 4,800.00", fb["amount_cents"] == 480_000, fb)
fc = pc_log.hire_fee(JOB, C)
check("no travel rate: travel days flagged, not paid", fc["missing"] == ["travel"] and fc["amount_cents"] == 315_000, fc)
check("a check-in day is unpaid and never 'missing'",
      pc_log.hire_fee({**JOB, "crew_dates": {"checkin": [["2026-10-10", "2026-10-10"]]}}, A) ==
      {"amount_cents": 0, "last_date": "2026-10-10", "detail": "", "missing": []})
ft = pc_log.hire_fee(JOB, T)
check("talent: their own days only, plus the agency fee: 1,200.00", ft["amount_cents"] == 120_000, ft)
check("and the line says so", ft["detail"] == "1 shoot x $1,000.00 + 20% agency fee", ft["detail"])
check("rates in dollars become exact cents", pc_log.to_cents(1300.5) == 130_050 and pc_log.to_cents("0.1") == 10
      and pc_log.to_cents(None) is None and pc_log.to_cents(-1) is None)

print("\nGenerated lines")
DAYS = [{"id": "d1", "name": "Day 1", "shoot_date": "2026-10-13"}]
LOCS = {"d1": {"project_id": "d1", "location_name": "Example Studio", "address": "2 Main St"}}
from services import per_diem, vendors as vendors_svc  # noqa: E402
hires = [A, B, C, T, HOLD, GONE]
resolved = vendors_svc.resolve(JOB, JOB["vendors"], DAYS, LOCS)
gen = pc_log.generated_lines(JOB, hires, resolved, per_diem.build_sheet(JOB, hires))
by_key = {g["key"]: g for g in gen}
check("one line per booked person; holds and released people aren't on it",
      [k for k in by_key if k.startswith("hire:")] == ["hire:h1", "hire:h2", "hire:h3", "hire:h4"])
check("a person's line is in their column", by_key["hire:h3"]["category"] == "stylist_hmu")
check("a missing rate travels with the line", by_key["hire:h3"]["missing_rates"] == ["travel"])
pd_total = per_diem.build_sheet(JOB, hires)["total"]
check("per diem is one Production/Misc line, the sheet's total",
      by_key["per_diem"]["amount_cents"] == pc_log.to_cents(pd_total) and by_key["per_diem"]["category"] == "production_misc")
check("per diem: 3 crew away (talent never get it): 450 + 380 + 450", by_key["per_diem"]["amount_cents"] == 128_000,
      by_key["per_diem"]["amount_cents"])
check("each vendor payment is its own line: hotel deposit and final bill",
      by_key["vendor:v1:deposit"]["amount_cents"] == 50_000 and by_key["vendor:v1:final_bill"]["amount_cents"] == 130_050)
check("the hotel is Production/Misc, named from its stay",
      by_key["vendor:v1:final_bill"]["category"] == "production_misc" and by_key["vendor:v1:final_bill"]["payee"] == "Example Inn")
check("the final bill is labelled as the balance", by_key["vendor:v1:final_bill"]["description"] == "Hotel - Final bill (balance)")
check("a location is Location/Studio", by_key["vendor:v2:final_bill"]["category"] == "location_studio")
check("no amount, no line (caterer, a $0 deposit); a broken link is left off",
      not any(k.startswith(("vendor:v3", "vendor:v4", "vendor:v5")) for k in by_key))

print("\nHand changes")
ov = {"hire:h2": {"amount_cents": 400_000}, "hire:h1": {"category": "video"}, "per_diem": {"hidden": True}}
rows = pc_log.apply_overrides(gen, ov)
r = {x["key"]: x for x in rows}
check("an amount is replaced and marked", r["hire:h2"]["amount_cents"] == 400_000 and r["hire:h2"]["amount_overridden"]
      and r["hire:h2"]["generated_amount_cents"] == 480_000)
check("a column is replaced and marked", r["hire:h1"]["category"] == "video" and r["hire:h1"]["category_overridden"])
check("a hidden line stays listed, marked hidden", r["per_diem"]["hidden"])
for bad in (("expense:x", {}), ("hire:h1", {"amount_cents": -1}), ("hire:h1", {"amount_cents": 1.5}),
            ("hire:h1", {"category": "snacks"}), ("drop table", {})):
    try:
        pc_log.validate_override(*bad)
        check(f"refused: {bad}", False, "accepted")
    except ValueError:
        check(f"refused: {bad}", True)
check("an empty change clears the entry", pc_log.validate_override("hire:h1", {"hidden": False}) is None)

print("\nThe log")
EXP = [{"id": "e1", "spent_on": "2026-10-11", "payee": "Example Air", "description": "Flights",
        "amount_cents": 84_260, "category": "production_misc"}]
INV = [{"amount_cents": 1_000_000, "paid": True}, {"amount_cents": 5, "paid": False}]
hdr = pc_log.header(JOB, ["2026-10-13"], "Alex Example")
log = pc_log.build_log(JOB, gen, EXP, INV, hdr)
t = log["totals"]
want_total = 540_000 + 480_000 + 315_000 + 120_000 + 128_000 + 50_000 + 130_050 + 240_000 + 84_260
check("total: every generated line and the receipt", t["total_cents"] == want_total == 2_087_310, t["total_cents"])
check("8% once on the total", t["fee_cents"] == 166_985 and t["grand_total_cents"] == 2_254_295)
check("less the paid advance", t["difference_cents"] == 2_254_295 - 1_000_000)
check("the header", hdr == {"job_name": "Riverbend Fall", "client": "Example Co", "producer": "Alex Example",
                            "project_dates": "Oct 11 - Oct 17, 2026"}, hdr)
nums = [x["number"] for x in log["rows"]]
check("numbered 1..n in date order", nums == list(range(1, len(nums) + 1)))
check("in date order, undated last", [x["spent_on"] for x in log["rows"]] == sorted(x["spent_on"] for x in log["rows"]))
hidden_log = pc_log.build_log({**JOB, "pc_log_generated": {"per_diem": {"hidden": True}}}, gen, EXP, INV, hdr)
check("a hidden line is off the total and unnumbered",
      hidden_log["totals"]["total_cents"] == want_total - 128_000
      and next(x for x in hidden_log["rows"] if x["key"] == "per_diem")["number"] is None)

print("\nMark as sent")
snap = pc_log.snapshot(hidden_log, "2026-10-21")
check("the copy holds only what was on the log", all(x["key"] != "per_diem" for x in snap["rows"]))
check("nothing has moved since", pc_log.changes_since(snap, hidden_log) == [])
A2 = {**A, "rate_shoot": 1300}
gen2 = pc_log.generated_lines(JOB, [A2, B, C, T], resolved, per_diem.build_sheet(JOB, [A2, B, C, T]))
later = pc_log.build_log({**JOB, "pc_log_generated": {"per_diem": {"hidden": True}}, "production_fee_bp": 1000},
                         gen2, EXP + [{"id": "e2", "spent_on": "2026-10-18", "payee": "Late cab", "description": "",
                                       "amount_cents": 4_500, "category": "production_misc"}], [], hdr)
ch = {c["key"]: c for c in pc_log.changes_since(snap, later)}
check("a rate change is flagged on that line, with both amounts",
      ch["hire:h1"]["change"] == "changed" and ch["hire:h1"]["then_cents"] == 540_000
      and ch["hire:h1"]["now_cents"] == 540_000 + 2 * 10_000 + 5_000, ch.get("hire:h1"))
check("a new receipt is flagged as added", ch["expense:e2"]["change"] == "added")
check("the fee % and the advance are flagged", "fee" in ch and "advance" in ch)
check("the copy itself is untouched", snap["totals"]["total_cents"] == want_total - 128_000)

print("\nThe files")
wb_bytes = to_xlsx(hidden_log)
from openpyxl import load_workbook  # noqa: E402
ws = load_workbook(io.BytesIO(wb_bytes)).active
check("her headings", [ws.cell(row=8, column=c).value for c in range(1, 6)] == ["#", "DATE", "TO WHOM PAID", "FOR WHAT", "TOTAL"])
check("her nine columns, in order", [ws.cell(row=8, column=c).value for c in range(6, 15)] == list(pc_log.CATEGORIES.values()))
check("JOB NAME at the top", ws["A1"].value == "JOB NAME:" and ws["B1"].value == "Riverbend Fall")
n_rows = sum(1 for x in hidden_log["rows"] if not x.get("hidden"))
tot = 9 + n_rows
check("each row's TOTAL is a live sum", ws["E9"].value == "=SUM(F9:N9)")
check("the footer is live formulas", ws.cell(row=tot, column=5).value == f"=SUM(E9:E{tot - 1})"
      and ws.cell(row=tot + 1, column=5).value == f"=ROUND(E{tot}*800/10000,2)"
      and ws.cell(row=tot + 4, column=5).value == f"=E{tot + 2}-E{tot + 3}")
check("the advance is a number", ws.cell(row=tot + 3, column=5).value == 10_000)
check("amounts sit in their category's column",
      ws.cell(row=9, column=6 + list(pc_log.CATEGORIES).index(hidden_log["rows"][0]["category"])).value
      == hidden_log["rows"][0]["amount_cents"] / 100)
html = to_html(hidden_log)
check("the PDF's own marks are ASCII only", html.isascii(), [ch for ch in html if not ch.isascii()][:5])
check("the PDF carries her footer", "LESS ADVANCE RECEIVED" in html and "DIFFERENCE BETWEEN ADVANCE AND ACTUALS" in html)
check("the PDF shows the grand total", pc_log.dollars(hidden_log["totals"]["grand_total_cents"]) in html)
check("the sent date prints when it's a sent copy", "Sent 10/21/26" in to_html(snap) and "Sent" not in html)

print("\nRoutes")
db = {"jobs": [dict(JOB)], "job_hires": [dict(h) for h in hires],
      "projects": [{"id": "d1", "job_id": J, "name": "Day 1", "shoot_date": "2026-10-13"}],
      "project_locations": [LOCS["d1"]]}
client = FakeClient(db)
for mod in (pc_router, budget_router, jobs_router, vendors_router):
    mod.get_client = lambda: client
pc_router._producer = lambda job: "Alex Example"
pc_router.render_call_sheet_pdf = lambda html: html.encode()
OWN = {"o1"}

v = pc_router.get_pc_log(J, owner_ids=OWN)
check("the live log loads, migrated, nothing sent", v["migrated"] and v["sent"] is None and v["changes"] == [])
check("with every generated line", sum(1 for x in v["rows"] if x["generated"]) == len(gen))
budget_router.add_expense(J, ExpenseCreate(payee="Example Air", description="Flights", amount_cents=84_260,
                                           spent_on="2026-10-11"), owner_ids=OWN)
budget_router.add_invoice(J, InvoiceCreate(label="Advance", amount_cents=1_000_000, paid=True), owner_ids=OWN)
v = pc_router.get_pc_log(J, owner_ids=OWN)
check("a typed receipt and the advance land on it", v["totals"]["total_cents"] == want_total
      and v["totals"]["advance_cents"] == 1_000_000, v["totals"])
v = pc_router.update_settings(J, pc_router.SettingsBody(production_fee_bp=1000, pc_log_camera="video"), owner_ids=OWN)
check("the fee % and the camera column are saved",
      v["settings"] == {"production_fee_bp": 1000, "pc_log_camera": "video"}
      and next(x for x in v["rows"] if x["key"] == "hire:h1")["category"] == "video")
pc_router.update_settings(J, pc_router.SettingsBody(production_fee_bp=800, pc_log_camera="photography"), owner_ids=OWN)
v = pc_router.update_generated(J, "per_diem", pc_router.GeneratedBody(hidden=True), owner_ids=OWN)
check("a generated line is hidden", v["totals"]["total_cents"] == want_total - 128_000)
v = pc_router.update_generated(J, "per_diem", pc_router.GeneratedBody(), owner_ids=OWN)
check("and put back", v["totals"]["total_cents"] == want_total and "per_diem" not in client.db["jobs"][0]["pc_log_generated"])
try:
    pc_router.update_generated(J, "hire:h1", pc_router.GeneratedBody(category="snacks"), owner_ids=OWN)
    check("a bad change is a 422", False, "accepted")
except HTTPException as e:
    check("a bad change is a 422", e.status_code == 422)

v = pc_router.mark_sent(J, pc_router.SendBody(sent_on="2026-10-21"), owner_ids=OWN)
check("Mark as sent freezes a copy", v["sent"]["sent_on"] == "2026-10-21" and len(v["sends"]) == 1 and v["changes"] == [])
client.db["job_hires"][0]["rate_shoot"] = 1300
v = pc_router.get_pc_log(J, owner_ids=OWN)
check("a rate changed afterwards is flagged, not applied to the copy",
      [c["key"] for c in v["changes"]] == ["hire:h1"] and v["totals"]["total_cents"] == want_total + 25_000)
xlsx = pc_router.download_xlsx(J, owner_ids=OWN)
ws = load_workbook(io.BytesIO(xlsx.body)).active
amounts = [ws.cell(row=r, column=c).value for r in range(9, 30) for c in range(6, 15)]
check("the download prints the copy that was sent (5,400.00, not 5,650.00)", 5400 in amounts and 5650 not in amounts)
check("named for the job", "PC%20Log%20-%20Riverbend%20Fall.xlsx" in xlsx.headers["content-disposition"])
pdf = pc_router.download_pdf(J, owner_ids=OWN)
check("the PDF prints the sent copy too", b"$5,400.00" in pdf.body and b"Sent 10/21/26" in pdf.body)
v = pc_router.mark_sent(J, pc_router.SendBody(sent_on="2026-10-25"), owner_ids=OWN)
check("sending again keeps both versions; the flag clears", len(v["sends"]) == 2 and v["changes"] == [])
first = v["sends"][0]["id"]
old = pc_router.download_xlsx(J, send_id=first, owner_ids=OWN)
ws = load_workbook(io.BytesIO(old.body)).active
check("any earlier version prints again", 5400 in [ws.cell(row=r, column=c).value for r in range(9, 30) for c in range(6, 15)])
try:
    pc_router.download_xlsx(J, send_id="nope", owner_ids=OWN)
    check("an unknown version is a 404", False, "no error")
except HTTPException as e:
    check("an unknown version is a 404", e.status_code == 404)

empty = FakeClient({"jobs": [{**JOB, "crew_dates": {}, "vendors": [], "id": "j2"}], "job_hires": [], "projects": []})
for mod in (pc_router, budget_router, jobs_router, vendors_router):
    mod.get_client = lambda: empty
try:
    pc_router.mark_sent("j2", pc_router.SendBody(), owner_ids=OWN)
    check("an empty log can't be sent", False, "sent")
except HTTPException as e:
    check("an empty log can't be sent", e.status_code == 422)

print("\nBefore supabase/reconciliation.sql")
pre_job = {k: v for k, v in JOB.items() if k not in ("production_fee_bp", "pc_log_camera", "pc_log_generated")}
pre = FakeClient({"jobs": [pre_job], "job_hires": [dict(h) for h in hires], "projects": db["projects"],
                  "project_locations": db["project_locations"]},
                 schema={"jobs": set(pre_job), "job_hires": set().union(*map(set, hires)),
                         "projects": {"id", "job_id", "name", "shoot_date"},
                         "project_locations": {"project_id", "location_name", "address"}})
for mod in (pc_router, budget_router, jobs_router, vendors_router):
    mod.get_client = lambda: pre
v = pc_router.get_pc_log(J, owner_ids=OWN)
check("the log still shows its generated lines, unmigrated",
      v["migrated"] is False and v["totals"]["fee_bp"] == 800 and any(x["generated"] for x in v["rows"]))
for name, fn in (("Mark as sent", lambda: pc_router.mark_sent(J, pc_router.SendBody(), owner_ids=OWN)),
                 ("hiding a line", lambda: pc_router.update_generated(J, "per_diem", pc_router.GeneratedBody(hidden=True), owner_ids=OWN)),
                 ("the fee %", lambda: pc_router.update_settings(J, pc_router.SettingsBody(production_fee_bp=500), owner_ids=OWN))):
    try:
        fn()
        check(f"{name} is a 503 naming the migration", False, "no error")
    except HTTPException as e:
        check(f"{name} is a 503 naming the migration", e.status_code == 503 and "reconciliation.sql" in e.detail, e.detail)

print()
if FAILS:
    print(f"FAILED: {len(FAILS)}")
    sys.exit(1)
print("All PC log checks passed.")
